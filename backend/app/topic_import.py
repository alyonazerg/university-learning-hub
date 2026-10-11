"""Preview spreadsheet topic updates before committing a whole group plan."""
import base64
import binascii
import csv
import io
from datetime import datetime, date, time, timezone
from pathlib import Path
from fastapi import APIRouter, Depends, HTTPException
from pydantic import Field, ValidationError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from .learning import StrictModel, administrator, database, no_store, group_for
from .models import Group, LessonPlan
from .planning import PlanInput, validate_schedule

router = APIRouter(prefix='/learning/topic-import', dependencies=[Depends(administrator), Depends(no_store)])
COLUMNS = {'position','topic','kind','date','time','duration_minutes','objectives','materials','test_url'}


class TopicImport(StrictModel):
    group_ids: list[int] = Field(min_length=1, max_length=30)
    filename: str = Field(max_length=200)
    content: str = Field(max_length=2800000)


def records(payload):
    try:
        raw = base64.b64decode(payload.content, validate=True)
        if len(raw)>2000000:
            raise ValueError('File must be under 2 MB')
        extension=Path(payload.filename).suffix.lower()
        if extension=='.csv':
            text=raw.decode('utf-8-sig')
            delimiter=';' if text.splitlines()[0].count(';')>text.splitlines()[0].count(',') else ','
            rows=list(csv.reader(io.StringIO(text), delimiter=delimiter))
        elif extension=='.xlsx':
            from openpyxl import load_workbook
            import zipfile
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(entry.file_size for entry in archive.infolist())>10000000:
                    raise ValueError('Workbook is too large; export the topic sheet as CSV')
            book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
            sheet=next((sheet for sheet in book if sheet.sheet_state=='visible'),None)
            if sheet is None: raise ValueError('No visible worksheet')
            if sheet.max_row>121 or sheet.max_column>9: raise ValueError('Use up to 120 topic rows and the documented columns')
            rows=list(sheet.iter_rows(max_row=122,max_col=10,values_only=True));book.close()
        elif extension=='.xls':
            import xlrd
            book=xlrd.open_workbook(file_contents=raw)
            sheet=next((sheet for sheet in book.sheets() if sheet.visibility==0),None)
            if sheet is None: raise ValueError('No visible worksheet')
            if sheet.nrows>121 or sheet.ncols>9: raise ValueError('Use up to 120 topic rows and the documented columns')
            rows=[]
            for index in range(sheet.nrows):
                row=[]
                for cell in sheet.row(index):
                    row.append(xlrd.xldate_as_datetime(cell.value,book.datemode) if cell.ctype==xlrd.XL_CELL_DATE else cell.value)
                rows.append(row)
        else: raise ValueError('Use CSV, XLS or XLSX')
        rows=[list(row) for row in rows if any(value not in ('',None) for value in row)]
        if not rows or len(rows)>121: raise ValueError('Use a header and 1–120 topic rows')
        while rows[0] and rows[0][-1] in ('',None): rows[0].pop()
        headers=[str(value).strip().lower() for value in rows[0]]
        if len(set(headers))!=len(headers) or not {'position','topic'}<=set(headers) or set(headers)-COLUMNS:
            raise ValueError('Required columns: position, topic. Optional: kind, date, time, duration_minutes, objectives, materials, test_url')
        data=[]
        for row in rows[1:]:
            if any(value not in ('',None) for value in row[len(headers):]): raise ValueError('Data outside named columns')
            data.append(dict(zip(headers,row)))
        if not data: raise ValueError('No topic rows')
        return data
    except (ValueError, UnicodeError, binascii.Error, csv.Error) as error:
        raise HTTPException(422,str(error))
    except Exception:
        raise HTTPException(422,'Cannot read this workbook. Use the provided template and save as CSV, XLS or XLSX.')


def prepare(payload,session,save):
    data=records(payload)
    if len(set(payload.group_ids))!=len(payload.group_ids): raise HTTPException(422,'Duplicate groups')
    summary=[]
    for gid in sorted(payload.group_ids):
        group_for(gid,session)
        group=session.scalar(select(Group).where(Group.id==gid).with_for_update())
        existing=list(session.scalars(select(LessonPlan).where(LessonPlan.group_id==gid)))
        positions={}; changes=[]; seen=set()
        for row in existing:
            positions.setdefault(row.position,[]).append(row)
        for index,record in enumerate(data,2):
            try:
                number=float(record['position'])
                if not number.is_integer(): raise ValueError('Topic number must be an integer')
                position=int(number)
                if position in seen: raise ValueError('Duplicate topic number')
                seen.add(position)
                matches=positions.get(position,[])
                if len(matches)>1: raise ValueError('Multiple existing lessons have this topic number; correct their numbering first')
                old=matches[0] if matches else None
                values={key:getattr(old,key) for key in PlanInput.model_fields} if old else {'group_id':gid,'position':position}
                if values.get('starts_at'):
                    values['starts_at']=values['starts_at'].replace(tzinfo=timezone.utc)
                values['topic']=str(record.get('topic') or '').strip()
                for key in ['kind','objectives','materials','duration_minutes','test_url']:
                    if record.get(key) not in ('',None): values[key]=record[key]
                if values.get('kind')=='test': values['duration_minutes']=90
                if record.get('date') not in ('',None):
                    day=record['date']
                    day=day.date().isoformat() if isinstance(day,datetime) else day.isoformat() if isinstance(day,date) else str(day).strip()
                    clock=record.get('time')
                    if clock in ('',None):
                        if old and old.starts_at: clock=(old.starts_at.hour+3)%24; clock=f'{clock:02d}:{old.starts_at.minute:02d}'
                        else: raise ValueError('Provide time together with date')
                    elif isinstance(clock,(datetime,time)): clock=clock.strftime('%H:%M')
                    values['starts_at']=datetime.fromisoformat(f'{day}T{clock}:00+03:00')
                elif record.get('time') not in ('',None): raise ValueError('Provide date together with time')
                parsed=PlanInput.model_validate(values)
                if old and old.attendance_lesson_id and (old.topic!=parsed.topic or old.starts_at!=parsed.starts_at):
                    raise ValueError('Topic/date is linked to an attendance register and cannot be replaced')
                changes.append((old,parsed))
            except (ValueError,TypeError,ValidationError) as error:
                raise HTTPException(422,f'Group {group.name}, row {index}: {error}')
        updated={old.id for old,_ in changes if old}
        candidates=[row for row in existing if row.id not in updated]+[parsed for _,parsed in changes]
        validate_schedule(candidates)
        summary.append({'group':group.name,'add':sum(old is None for old,_ in changes),'update':sum(old is not None for old,_ in changes),'topics':[{'position':row.position,'topic':row.topic,'kind':row.kind,'starts_at':row.starts_at,'duration_minutes':row.duration_minutes,'test_url':row.test_url} for _,row in changes]})
        if save:
            for old,parsed in changes:
                if old:
                    for key,value in parsed.model_dump().items(): setattr(old,key,value)
                else: session.add(LessonPlan(**parsed.model_dump()))
    return {'groups':summary}


@router.post('/preview')
def preview(payload:TopicImport,session=Depends(database)):
    return prepare(payload,session,False)


@router.post('/apply')
def apply(payload:TopicImport,session=Depends(database)):
    try:
        result=prepare(payload,session,True);session.commit();return result
    except (HTTPException,IntegrityError) as error:
        session.rollback()
        if isinstance(error,IntegrityError): raise HTTPException(409,'Lessons overlap or the plan changed; preview again')
        raise
