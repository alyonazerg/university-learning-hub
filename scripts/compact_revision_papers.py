"""Compact existing revision papers without modifying questions, choices or essay topics.
Run with python-docx installed; rendered page count must be verified separately.
"""
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml.ns import qn
import re,json
root=Path(__file__).resolve().parents[1]/'frontend/templates';before={}
for p in sorted(root.glob('Revision-Unit-7-Group-*.docx')):
 if not any(p.name.endswith(suffix) for suffix in ['-ZipGrade.docx','-paper.docx','-With-Reading-Preview.docx']):continue
 group=p.name;d=Document(p);before[group]=[p.text for p in d.paragraphs if p.text]
 for sec in d.sections:
  sec.top_margin=sec.bottom_margin=Cm(1.15);sec.left_margin=sec.right_margin=Cm(1.25)
  sec.header_distance=sec.footer_distance=Cm(.5)
 normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1;normal.paragraph_format.space_after=Pt(2)
 for name,size in [('Title',15),('Heading 1',11.5),('Heading 2',11)]:
  st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.paragraph_format.space_before=Pt(6);st.paragraph_format.space_after=Pt(3)
 questions=False
 for para in list(d.paragraphs):
  text=para.text
  if para.style.name=='Title' and text.startswith('RT1'):questions=True
  # Keep the separate preview page, remove only redundant section page breaks.
  if not text and questions and para._p.xpath('.//w:br[@w:type="page"]'):
   para._element.getparent().remove(para._element);continue
  fmt=para.paragraph_format;fmt.line_spacing=1;fmt.space_before=Pt(0);fmt.space_after=Pt(2);fmt.keep_together=True;fmt.widow_control=True
  if para.style.name=='Title':fmt.space_after=Pt(5)
  elif para.style.name.startswith('Heading'):fmt.space_before=Pt(6);fmt.space_after=Pt(3);fmt.keep_with_next=True
  for run in para.runs:
   run.font.name='Arial'
   if run.font.size:run.font.size=Pt(11)
  if questions and re.match(r'^\d+\. ',text) and 'Essay' not in text:
   # Numbered prompts and options are kept as one unit, except essay topics.
   fmt.keep_with_next=not any(text.startswith(f'{i}. '+prefix) for i in range(1,5) for prefix in ['Our roots','The benefits','Traditions','A sense','Meaningful','Having','Companies','Paying'])
   if fmt.keep_with_next:
    for run in para.runs:run.bold=True
   fmt.space_before=Pt(5)
  if re.match(r'^[A-E]\. ',text):fmt.keep_with_next=not text.startswith('E. ');fmt.space_after=Pt(0)
 # Place short alternatives in two aligned columns without shrinking the type.
 paragraphs=list(d.paragraphs)
 for index,para in enumerate(paragraphs):
  if not re.match(r'^[AC]\. ',para.text) or index+1>=len(paragraphs):continue
  following=paragraphs[index+1]
  if not following.text.startswith('B. ' if para.text.startswith('A. ') else 'D. '):continue
  if max(len(para.text),len(following.text))>40:continue
  para.add_run('\t'+following.text);para.paragraph_format.tab_stops.add_tab_stop(Cm(9.1))
  following._element.getparent().remove(following._element)
 d.save(p)
 normalize=lambda value:re.sub(r'\s+',' ',value).strip()
 assert normalize(' '.join(p.text for p in d.paragraphs if p.text))==normalize(' '.join(before[group]))
print(f'Compacted {len(before)} question documents; text preserved.')
